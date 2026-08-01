# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(5.625)

BG_DARK = RGBColor(0x0D, 0x11, 0x17)
BG_LIGHT = RGBColor(0xF5, 0xF5, 0xF7)
PRIMARY = RGBColor(0x00, 0x7A, 0xCC)
ACCENT = RGBColor(0x00, 0xB8, 0xD4)
TEXT_LIGHT = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_DARK = RGBColor(0x1A, 0x1A, 0x2E)
TEXT_MUTED = RGBColor(0x6B, 0x7B, 0x8D)
CARD_BG = RGBColor(0xE8, 0xEE, 0xF4)
GREEN = RGBColor(0x00, 0xC8, 0x53)
ORANGE = RGBColor(0xFF, 0x98, 0x00)
RED = RGBColor(0xFF, 0x17, 0x44)
NAVY = RGBColor(0x1A, 0x23, 0x7E)
TEAL = RGBColor(0x00, 0x96, 0x88)

def dark_slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG_DARK
    return s

def light_slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG_LIGHT
    return s

def txt(s, text, l, t, w, h, sz=14, c=TEXT_DARK, b=False, a=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(sz)
    p.font.bold = b
    p.font.color.rgb = c
    p.alignment = a
    return tb

def bullets(s, items, l, t, w, h, sz=12, c=TEXT_DARK, sp=Pt(4)):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(sz)
        p.font.color.rgb = c
        p.space_after = sp
    return tb

def card(s, l, t, w, h, fc=CARD_BG):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fc
    sh.line.fill.background()
    return sh

def pill(s, l, t, w, h, text, fc=PRIMARY, tc=TEXT_LIGHT, sz=11):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fc
    sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(sz)
    p.font.color.rgb = tc
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    return sh

def make_table(s, headers, rows, l, t, w, h, col_widths, hdr_color=PRIMARY):
    tbl = s.shapes.add_table(len(rows)+1, len(headers), Inches(l), Inches(t), Inches(w), Inches(h)).table
    for ci, cw in enumerate(col_widths):
        tbl.columns[ci].width = Inches(cw)
    for ci, hd in enumerate(headers):
        c = tbl.cell(0, ci)
        c.text = hd
        for p in c.text_frame.paragraphs:
            p.font.bold = True; p.font.size = Pt(10); p.font.color.rgb = TEXT_LIGHT
        c.fill.solid(); c.fill.fore_color.rgb = hdr_color
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.cell(ri+1, ci)
            c.text = val
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(9); p.font.color.rgb = TEXT_DARK

# ===== 1. Cover =====
s = dark_slide()
txt(s, "基于CH32H417双核MCU的", 0.6, 0.8, 8.8, 0.5, sz=18, c=ACCENT, a=PP_ALIGN.CENTER)
txt(s, "EEG脑电采集与运动想象方向识别系统", 0.6, 1.4, 8.8, 0.8, sz=34, c=TEXT_LIGHT, b=True, a=PP_ALIGN.CENTER)
txt(s, "技术验证平台", 0.6, 2.4, 8.8, 0.5, sz=20, c=ORANGE, b=True, a=PP_ALIGN.CENTER)
txt(s, "验证从采集→处理→推理→显示的完整BCI链路可行性", 0.6, 2.9, 8.8, 0.4, sz=13, c=TEXT_MUTED, a=PP_ALIGN.CENTER)
card(s, 1.5, 3.5, 7.0, 0.9, fc=RGBColor(0x15, 0x1C, 0x25))
txt(s, "V3F 100MHz + V5F 400MHz  |  ADS1299 8ch 24bit  |  ICM-42605 IMU  |  ESP8266 WiFi", 1.8, 3.65, 6.4, 0.6, sz=12, c=ACCENT, a=PP_ALIGN.CENTER)

# ===== 2. Architecture =====
s = light_slide()
txt(s, "系统架构：为什么是三层？", 0.6, 0.25, 8.8, 0.6, sz=28, c=PRIMARY, b=True)

card(s, 0.3, 1.0, 2.9, 2.4, fc=RGBColor(0xE3, 0xF2, 0xFD))
txt(s, "采集端 (MCU)", 0.5, 1.1, 2.5, 0.3, sz=13, c=PRIMARY, b=True)
bullets(s, [
    "为什么不用纯Android?",
    "→ ADS1299需SPI+DMA硬实时",
    "→ 250SPS×8ch Android无法保证",
    "→ FFT需确定性延迟",
], 0.5, 1.45, 2.5, 1.8, sz=10)

card(s, 3.5, 1.0, 2.9, 2.4, fc=RGBColor(0xE8, 0xF5, 0xE9))
txt(s, "医生端 (Android)", 3.7, 1.1, 2.5, 0.3, sz=13, c=GREEN, b=True)
bullets(s, [
    "为什么不能省?",
    "→ 患者无法自行操作设备",
    "→ 需专业监控+训练配置",
    "→ 姿态警报需医护响应",
], 3.7, 1.45, 2.5, 1.8, sz=10)

card(s, 6.7, 1.0, 3.0, 2.4, fc=RGBColor(0xFD, 0xE0, 0xDC))
txt(s, "患者端 (Android)", 6.9, 1.1, 2.6, 0.3, sz=13, c=ORANGE, b=True)
bullets(s, [
    "为什么独立?",
    "→ AR眼镜投屏需独立渲染",
    "→ SSVEP闪烁需精确帧率",
    "→ 被动化避免误操作",
], 6.9, 1.45, 2.6, 1.8, sz=10)

txt(s, "通信端口", 0.3, 3.6, 2.0, 0.3, sz=12, c=PRIMARY, b=True)
make_table(s,
    ["端口", "方向", "用途"],
    [
        ("41002", "MCU ↔ 医生端", "双向透传: 命令/事件"),
        ("41003", "医生端 → Python", "原始帧转发(SSVEP分析)"),
        ("41004", "医生端 ↔ 患者端", "控制: PAGE/SSVEP/TARGET + 握手心跳"),
        ("41005", "UDP广播", "患者端自动发现医生端IP"),
    ],
    0.3, 3.95, 9.4, 1.4, [1.0, 2.5, 5.9])

# ===== 3. Dual Core =====
s = light_slide()
txt(s, "双核异构：为什么不用单核？", 0.6, 0.25, 8.8, 0.6, sz=28, c=PRIMARY, b=True)

card(s, 0.3, 1.0, 4.4, 2.4, fc=RGBColor(0xE3, 0xF2, 0xFD))
txt(s, "V3F 100MHz — 采集与处理", 0.5, 1.1, 4.0, 0.3, sz=13, c=RGBColor(0x00, 0x5A, 0x9E), b=True)
bullets(s, [
    "ADS1299 SPI3+DMA 8ch 250SPS采集",
    "50Hz陷波 + 带通滤波 + 基线去除",
    "256点FFT (步进128, 50%重叠, 500ms)",
    "8ch频谱分片发送(32片×16字节)",
    "6ch加权专注度 + ICM-42605姿态",
    "DISPLAY_CFG按需配置发送通道",
], 0.5, 1.45, 4.0, 1.8, sz=10)

card(s, 5.3, 1.0, 4.4, 2.4, fc=RGBColor(0xE8, 0xF5, 0xE9))
txt(s, "V5F 400MHz — 推理", 5.5, 1.1, 4.0, 0.3, sz=13, c=GREEN, b=True)
bullets(s, [
    "IPC共享内存接收V3F预处理帧",
    "24维频域特征 → LEFT/RIGHT分类",
    "500ms推理1次, 4次投票, 2s出结果",
    "支持FFT24/CSP双模型切换",
], 5.5, 1.45, 4.0, 1.8, sz=10)

card(s, 0.3, 3.6, 9.4, 1.7, fc=RGBColor(0xFD, 0xE0, 0xDC))
txt(s, "为什么不用单核?", 0.5, 3.7, 4.0, 0.3, sz=13, c=ORANGE, b=True)
bullets(s, [
    "采集+推理串行: 250SPS采集不允许被推理阻塞, 推理(矩阵运算)又需连续CPU时间",
    "→ 采集抖动→FFT窗口不对齐→频谱泄漏; 推理被采集中断打断→延迟翻倍",
    "双核: V3F保证采集确定性, V5F专注推理算力, IPC零拷贝并行互不阻塞",
], 0.5, 4.05, 9.0, 1.1, sz=10)

# ===== 4. Channel & Attention =====
s = light_slide()
txt(s, "通道选择与专注度", 0.6, 0.25, 8.8, 0.6, sz=28, c=PRIMARY, b=True)

card(s, 0.3, 1.0, 5.8, 2.1, fc=RGBColor(0xE3, 0xF2, 0xFD))
txt(s, "8通道布局 (10-20系统)", 0.5, 1.1, 5.4, 0.3, sz=13, c=PRIMARY, b=True)
make_table(s,
    ["通道", "脑区", "角色"],
    [
        ("F3 / F4", "前额叶", "执行控制/注意 → 专注度(0.18)"),
        ("CP3 / CP4", "中央-顶叶", "感觉运动整合 → 专注度(0.19) + MI主特征"),
        ("C3 / C4", "初级运动皮层", "运动执行 → 专注度(0.13, 运动伪迹大)"),
        ("O1 / OZ", "枕叶视觉区", "SSVEP稳态视觉诱发电位 → 权重0"),
    ],
    0.5, 1.45, 5.4, 1.5, [0.8, 1.2, 3.4])

card(s, 6.4, 1.0, 3.3, 2.1, fc=RGBColor(0xFD, 0xE0, 0xDC))
txt(s, "权重依据", 6.6, 1.1, 2.9, 0.3, sz=13, c=ORANGE, b=True)
bullets(s, [
    "CP最高: MI核心区, μ/β节律最显著",
    "F次高: 前额叶与注意力强相关(文献共识)",
    "C最低: 初级运动皮层, 运动伪迹干扰大",
    "O=0: 视觉区对MI无贡献, 专用于SSVEP",
    "排序有神经生理学依据, 数值为经验设定",
], 6.6, 1.45, 2.9, 1.5, sz=9)

card(s, 0.3, 3.3, 4.5, 2.0, fc=RGBColor(0xE8, 0xF5, 0xE9))
txt(s, "专注度算法", 0.5, 3.4, 4.0, 0.3, sz=13, c=GREEN, b=True)
bullets(s, [
    "6通道α/β能量比 → 加权求和 → EMA平滑",
    "双状态: 瞬时(短窗) + 趋势(长窗EMA)",
    "500ms更新, 阈值可调适配不同基线",
], 0.5, 3.75, 4.0, 1.4, sz=10)

card(s, 5.0, 3.3, 4.7, 2.0, fc=RGBColor(0xFD, 0xE0, 0xDC))
txt(s, "已知局限", 5.2, 3.4, 4.3, 0.3, sz=13, c=ORANGE, b=True)
bullets(s, [
    "权重为经验设定, 非数据驱动优化",
    "→ 改进: 标注数据 + 回归/网格搜索",
    "未做受试者间差异分析",
    "→ 改进: 个性化权重校准流程",
], 5.2, 3.75, 4.3, 1.4, sz=10)

# ===== 5. Pipeline + Protocol =====
s = light_slide()
txt(s, "信号处理与通信协议", 0.6, 0.25, 8.8, 0.6, sz=28, c=PRIMARY, b=True)

steps = [
    ("采集", "ADS1299\n8ch×250SPS\n24bit", NAVY),
    ("滤波", "50Hz陷波\n带通滤波\n去基线", RGBColor(0x00, 0x5A, 0x9E)),
    ("FFT", "256点\n步进128\n50%重叠", PRIMARY),
    ("特征", "频谱+专注度\n+24维方向", TEAL),
    ("推理", "FFT24\n4次投票\n2s/结果", GREEN),
]
for i, (t, d, c) in enumerate(steps):
    x = 0.3 + i * 1.95
    pill(s, x, 0.9, 1.75, 0.4, t, fc=c, sz=12)
    card(s, x, 1.35, 1.75, 0.85, fc=CARD_BG)
    txt(s, d, x+0.1, 1.42, 1.55, 0.75, sz=9, c=TEXT_DARK, a=PP_ALIGN.CENTER)

txt(s, "二进制帧  0x7E帧头尾 + 0x7D转义", 0.3, 2.35, 5.0, 0.25, sz=11, c=PRIMARY, b=True)
make_table(s,
    ["CMD", "值", "负载", "频率"],
    [
        ("RAW/FILT_WAVE", "0x04/0x10", "[ch:1B][val:4B]", "250SPS/ch"),
        ("FOCUS", "0x05", "[attn0~ema1:16B][trend:1B][inst:1B]", "500ms"),
        ("SPECTRUM", "0x20+x", "[frag:1B][total:1B][float×4]×32片", "500ms"),
    ],
    0.3, 2.65, 9.4, 1.1, [1.8, 0.9, 5.1, 1.6])

txt(s, "文本行  \\r\\n结尾, seq序号去重", 0.3, 3.85, 5.0, 0.25, sz=11, c=RGBColor(0x00, 0x5A, 0x9E), b=True)
make_table(s,
    ["协议", "示例", "触发"],
    [
        ("RESULT", "RESULT,seq=5,INTENT=LEFT,S_LEFT=6500,CONF=3000", "~2s"),
        ("TASK", "TASK,LEFT,start / TASK,DONE,seq=5", "事件"),
        ("FALL", "FALL_EVENT,seq=1,gyro=35000,acc=250", "事件"),
    ],
    0.3, 4.15, 9.4, 1.0, [1.5, 6.5, 1.4], hdr_color=RGBColor(0x00, 0x5A, 0x9E))

txt(s, "频谱128点分32片避免WiFi丢包  |  seq序号防重复处理  |  带宽合计约8.4KB/s", 0.3, 5.25, 9.4, 0.3, sz=10, c=TEXT_MUTED, a=PP_ALIGN.CENTER)

# ===== 6. Android Dual App =====
s = light_slide()
txt(s, "Android双端", 0.6, 0.25, 8.8, 0.6, sz=28, c=PRIMARY, b=True)

card(s, 0.3, 1.0, 4.6, 4.2, fc=RGBColor(0xE8, 0xF5, 0xE9))
txt(s, "医生端 — 监控与决策", 0.5, 1.1, 4.2, 0.3, sz=13, c=GREEN, b=True)
bullets(s, [
    "8ch实时波形/频谱 (4通道同时监控)",
    "专注度双曲线 + 填充 + 瞬时/趋势状态",
    "MI训练: 序列控制 + DIRCSV特征采集",
    "方向推理: 循环/随机目标 + 自动验证",
    "  → TARGET命令下发患者端",
    "SSVEP: 4频率刺激控制 + 频谱对比",
    "姿态监护: 事件日志 + 坠床红色警报",
    "  → 久未翻身橙色提醒",
], 0.5, 1.45, 4.2, 3.5, sz=10)

card(s, 5.3, 1.0, 4.6, 4.2, fc=RGBColor(0xFD, 0xE0, 0xDC))
txt(s, "患者端 — 完全被动化", 5.5, 1.1, 4.2, 0.3, sz=13, c=ORANGE, b=True)
bullets(s, [
    "患者(脑卒中/运动障碍)无法操作手机",
    "→ 无TabLayout, 禁止手动翻页",
    "→ 所有页面由医生端PAGE命令驱动",
    "MI训练: 方向箭头引导 (◀/▶)",
    "MI推理: 目标+结果双箭头",
    "  → 一致绿色, 不一致红色, 1.5s消失",
    "SSVEP: 全屏闪烁 (11/13/15/17Hz)",
    "AR眼镜适配: 纯黑背景+200sp字号",
    "41004双向握手+PING/PONG+断线重连",
], 5.5, 1.45, 4.2, 3.5, sz=10)

# ===== 7. Honest Assessment =====
s = dark_slide()
txt(s, "技术亮点与现状评估", 0.6, 0.25, 8.8, 0.6, sz=28, c=ACCENT, b=True)

card(s, 0.3, 1.0, 4.6, 2.0, fc=RGBColor(0x15, 0x1C, 0x25))
txt(s, "工程亮点", 0.5, 1.1, 4.2, 0.3, sz=13, c=ACCENT, b=True)
bullets(s, [
    "双核异构: V3F采集+V5F推理并行, IPC零拷贝",
    "DISPLAY_CFG按需配置: 节省MCU CPU和带宽",
    "频谱分片: 128点分32片, 避免WiFi丢包",
    "患者端被动化: 命令驱动, 适配AR眼镜",
], 0.5, 1.4, 4.2, 1.4, sz=10, c=TEXT_MUTED)

card(s, 5.3, 1.0, 4.6, 2.0, fc=RGBColor(0x15, 0x1C, 0x25))
txt(s, "当前局限 (诚实评估)", 5.5, 1.1, 4.2, 0.3, sz=13, c=ORANGE, b=True)
bullets(s, [
    "专注度权重为经验设定, 非数据驱动",
    "推理模型受试者少, 泛化性未验证",
    "WiFi热点模式延迟波动, 非硬实时",
    "无电极阻抗监测, 接触不良无法检测",
], 5.5, 1.4, 4.2, 1.4, sz=10, c=TEXT_MUTED)

card(s, 0.3, 3.2, 9.6, 2.1, fc=RGBColor(0x15, 0x1C, 0x25))
txt(s, "项目定位", 0.5, 3.3, 9.2, 0.3, sz=13, c=TEXT_LIGHT, b=True)
bullets(s, [
    "这是一个技术验证平台, 不是临床产品 — 验证了从采集→处理→推理→显示的完整链路可行性",
    "离临床使用还有本质差距: 需要临床试验(伦理审批)、模型大量受试者训练、医疗器械认证(NMPA)",
    "当前价值: 搭通了嵌入式+Android全链路, 暴露了真实工程问题(WiFi延迟/信号质量/模型泛化), 明确了改进方向",
    "改进路线: BLE替代WiFi(低延迟) + 标注数据优化权重 + 电极阻抗监测 + 便携化(锂电池)",
], 0.5, 3.65, 9.2, 1.5, sz=10, c=TEXT_MUTED)

# ===== 8. End =====
s = dark_slide()
txt(s, "谢谢", 0.6, 1.8, 8.8, 1.0, sz=48, c=TEXT_LIGHT, b=True, a=PP_ALIGN.CENTER)
txt(s, "Q & A", 0.6, 3.0, 8.8, 0.6, sz=24, c=ACCENT, a=PP_ALIGN.CENTER)

output_path = r"D:\Libraries\Projects\EEG\output\EEG项目介绍.pptx"
prs.save(output_path)
print(f"Saved to {output_path}")
